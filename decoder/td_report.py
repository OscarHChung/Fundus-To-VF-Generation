"""Re-express the frozen OOF results (tag p1disc_denoise) in TOTAL DEVIATION (per point) and
MEAN DEVIATION (per eye), alongside raw sensitivity, to answer: "are raw-sensitivity and TD the
same for our purposes?"

Self-contained: numpy + openpyxl + matplotlib only. No torch, no model reload. Inputs:
  - decoder/results/auto/oof_cache_notta/p1disc_denoise_f{0..4}.npz   (vp/vt = held-out OOF preds/truths, 52-pt query order)
  - decoder/results/cv_long/fold{f}_val.json                          (row-aligned records -> PatientID/Laterality/VisitNumber)
  - data/vf_tests/uwhvf_vf_tests.json                                 (hvf & td -> reconstruct the 24-2 age-normative, exact)
  - data/vf_tests/grape_data.xlsx                                     (per-visit age = baseline age + cumulative interval)

TD(x)   = sensitivity(x) - normal(age, x)      [per point]
MD(eye) = mean_x TD(x)                          [per eye; unweighted approximation of HFA's variance-weighted MD]

  python decoder/td_report.py
"""
import os, sys, json
import numpy as np
import openpyxl
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AUTO = os.path.join(ROOT, "decoder", "results", "auto")
CV   = os.path.join(ROOT, "decoder", "results", "cv_long")
TAG  = "p1disc_denoise"

# ---- 24-2 grid (inlined from diagnostics.py; no torch import) ----
mask_OD = np.array([
    [False, False, False, True,  True,  True,  True,  False, False],
    [False, False, True,  True,  True,  True,  True,  True,  False],
    [False, True,  True,  True,  True,  True,  True,  True,  True ],
    [True,  True,  True,  True,  True,  True,  True,  False, True ],
    [True,  True,  True,  True,  True,  True,  True,  False, True ],
    [False, True,  True,  True,  True,  True,  True,  True,  True ],
    [False, False, True,  True,  True,  True,  True,  True,  False],
    [False, False, False, True,  True,  True,  True,  False, False]], dtype=bool)
valid_od = [i for i, v in enumerate(mask_OD.flatten()) if v]        # 52 flat-grid indices, query order
MASKED = 99.0


def pooled(preds, trues):
    """Pointwise pooled metrics over non-masked points. preds/trues: (N,52) with nan at masked."""
    m = ~np.isnan(trues)
    P = preds[m]; T = trues[m]
    d = P - T
    r = float(np.corrcoef(T, P)[0, 1])
    slope = float(np.cov(T, P, bias=True)[0, 1] / np.var(T))     # OLS slope of pred on true
    return dict(n=int(m.sum()), mae=float(np.mean(np.abs(d))), rmse=float(np.sqrt(np.mean(d**2))),
                bias=float(np.mean(d)), r=r, slope=slope)


def eye_metrics(mp, mt):
    """Per-eye index (MD or mean-sens) metrics. mp/mt: (N,) arrays."""
    d = mp - mt
    return dict(n=len(mt), mae=float(np.mean(np.abs(d))), bias=float(np.mean(d)),
                r=float(np.corrcoef(mt, mp)[0, 1]))


# ======================================================================================
# 1. Reconstruct the 24-2 age-normative from UWHVF:  normal(age,cell) = hvf - td
# ======================================================================================
print("[1] reconstructing 24-2 age-normative from UWHVF (normal = hvf - td) ...", flush=True)
uw = json.load(open(os.path.join(ROOT, "data", "vf_tests", "uwhvf_vf_tests.json")))
# accumulate per grid-cell (only the 52 query cells): lists of (age, normal)
acc_age = {c: [] for c in valid_od}
acc_nrm = {c: [] for c in valid_od}
for pid, rec in uw["data"].items():
    for side in ("L", "R"):
        for v in rec.get(side, []) or []:
            age = v["age"]
            hvf = np.asarray(v["hvf"], float).flatten()
            td  = np.asarray(v["td"],  float).flatten()
            for c in valid_od:
                h, t = hvf[c], td[c]
                if h < 90 and abs(t) < 90:          # skip blind-spot placeholder / masked
                    acc_age[c].append(age); acc_nrm[c].append(h - t)
# per-cell linear fit normal = a*age + b  -> 52-vectors a_slot, b_slot (query order)
a_slot = np.zeros(52); b_slot = np.zeros(52); resid = []
for j, c in enumerate(valid_od):
    A = np.array(acc_age[c]); N = np.array(acc_nrm[c])
    a, b = np.polyfit(A, N, 1)
    a_slot[j], b_slot[j] = a, b
    resid.append(np.std(N - (a * A + b)))
print(f"    per-cell age-fit residual: mean {np.mean(resid):.4f} dB, max {np.max(resid):.4f} dB "
      f"(->deterministic normative recovered)")
print(f"    normal@age50 range across 52 pts: {np.min(a_slot*50+b_slot):.1f}..{np.max(a_slot*50+b_slot):.1f} dB "
      f"(location spread = why r/slope move); age slope mean {np.mean(a_slot):.4f} dB/yr")


def normal52(age):
    return a_slot * age + b_slot


# ======================================================================================
# 2. Per-visit age from GRAPE (baseline age + cumulative interval)
# ======================================================================================
print("[2] building per-visit age map from GRAPE ...", flush=True)
wb = openpyxl.load_workbook(os.path.join(ROOT, "data", "vf_tests", "grape_data.xlsx"), read_only=True)
bs = list(wb["Baseline"].iter_rows(values_only=True)); bh = bs[0]
iS, iL, iA = bh.index("Subject Number"), bh.index("Laterality"), bh.index("Age")
base_age = {}
for r in bs[1:]:
    if r[iS] is None or not isinstance(r[iA], (int, float)):
        continue
    base_age[(int(r[iS]), str(r[iL]).strip()[:2])] = float(r[iA])
fu = list(wb["Follow-up"].iter_rows(values_only=True)); fh = fu[0]
jS, jL, jV, jI = fh.index("Subject Number"), fh.index("Laterality"), fh.index("Visit Number"), fh.index("Interval Years")
cum_int = {}
for r in fu[1:]:
    if r[jS] is None or r[jV] is None:
        continue
    cum_int[(int(r[jS]), str(r[jL]).strip()[:2], int(r[jV]))] = float(r[jI] or 0.0)
cohort_mean = float(np.mean(list(base_age.values())))


def age_of(pid, lat, visit):
    lat = str(lat).strip()[:2]
    base = base_age.get((pid, lat), cohort_mean)
    return base + cum_int.get((pid, lat, visit), 0.0)


# ======================================================================================
# 3. Load OOF cache + row-aligned records; convert to TD; collect eye-level MD
# ======================================================================================
print("[3] converting frozen OOF preds/truths to TD + MD ...", flush=True)
sens_p, sens_t, td_p, td_t = [], [], [], []
md_p, md_t, msens_p, msens_t = [], [], [], []
n_fallback = 0
for f in range(5):
    z = np.load(os.path.join(AUTO, "oof_cache_notta", f"{TAG}_f{f}.npz"))
    vp, vt = z["vp"], z["vt"]                       # (Nval,52) query order, masked truth >= 99
    recs = json.load(open(os.path.join(CV, f"fold{f}_val.json")))
    assert len(recs) == len(vp), (f, len(recs), len(vp))
    for i, rec in enumerate(recs):
        pid, lat, vis = int(rec["PatientID"]), rec["Laterality"], int(rec["VisitNumber"])
        if (pid, str(lat).strip()[:2]) not in base_age:
            n_fallback += 1
        age = age_of(pid, lat, vis)
        nrm = normal52(age)
        ps = vp[i].astype(float).copy(); ts = vt[i].astype(float).copy()
        ts[ts >= MASKED] = np.nan
        # mask preds wherever truth is masked (matches pooled_stats discipline)
        pt = np.where(np.isnan(ts), np.nan, ps - nrm)
        tt = ts - nrm
        sens_p.append(ps); sens_t.append(ts); td_p.append(pt); td_t.append(tt)
        # per-eye indices (over that eye's valid points)
        vmask = ~np.isnan(ts)
        md_p.append(float(np.mean((ps - nrm)[vmask])));   md_t.append(float(np.mean(tt[vmask])))
        msens_p.append(float(np.mean(ps[vmask])));        msens_t.append(float(np.mean(ts[vmask])))

sens_p = np.array(sens_p); sens_t = np.array(sens_t)
td_p = np.array(td_p); td_t = np.array(td_t)
md_p = np.array(md_p); md_t = np.array(md_t)
msens_p = np.array(msens_p); msens_t = np.array(msens_t)
print(f"    {len(sens_t)} eyes, {int((~np.isnan(sens_t)).sum())} points; age fallbacks: {n_fallback}")

# ======================================================================================
# 4. Report
# ======================================================================================
S = pooled(sens_p, sens_t)          # raw sensitivity space
D = pooled(td_p, td_t)              # total-deviation space
print("\n" + "=" * 78)
print("POINTWISE (pooled over all points)         RAW SENSITIVITY      TOTAL DEVIATION")
print("-" * 78)
for key, name in [("mae", "MAE            (dB)"), ("rmse", "RMSE           (dB)"),
                  ("bias", "bias  pred-true(dB)"), ("r", "Pearson r          "),
                  ("slope", "OLS slope pred~true")]:
    print(f"  {name:24s}   {S[key]:12.4f}     {D[key]:16.4f}")
print("-" * 78)
print("  MAE/RMSE/bias identical (normative cancels per point) | r & slope MOVE (TD removes")
print("  the location-normal gradient, i.e. the 'easy' between-location variance).")

MDs = eye_metrics(msens_p, msens_t)   # per-eye mean sensitivity
MDd = eye_metrics(md_p, md_t)         # per-eye MEAN DEVIATION
print("\n" + "=" * 78)
print("PER-EYE GLOBAL INDEX (n={} eyes)      MEAN SENSITIVITY        MEAN DEVIATION".format(len(md_t)))
print("-" * 78)
for key, name in [("mae", "index MAE      (dB)"), ("bias", "index bias     (dB)"),
                  ("r", "Pearson r          ")]:
    print(f"  {name:24s}   {MDs[key]:12.4f}     {MDd[key]:16.4f}")
print("-" * 78)
print("  MD-MAE == mean-sens MAE (per-eye normative is a per-eye constant -> cancels);")
print("  MD Pearson r differs slightly (between-eye normal spread from age).")

# severity strata by TRUE MD (Hodapp-Anderson-style), pointwise MAE within each (invariant to space)
print("\n" + "=" * 78)
print("SEVERITY STRATA cut on TRUE MEAN DEVIATION (Hodapp-Anderson)")
print("-" * 78)
bands = [("early   MD > -6",   md_t > -6),
         ("moderate -12..-6",  (md_t <= -6) & (md_t >= -12)),
         ("severe  MD < -12",  md_t < -12)]
for label, sel in bands:
    if sel.sum() == 0:
        print(f"  {label:20s}: 0 eyes"); continue
    st = sens_t[sel]; sp = sens_p[sel]
    mm = ~np.isnan(st)
    mae = float(np.mean(np.abs((sp - st)[mm])))
    print(f"  {label:20s}: {int(sel.sum()):3d} eyes   pointwise MAE {mae:5.2f} dB   "
          f"(true MD {md_t[sel].mean():6.2f})")
print("  (pointwise MAE within a fixed eye-set is the SAME number in raw or TD space; only the")
print("   BAND MEMBERSHIP is defined in MD space -- that is the part that isn't raw-sensitivity.)")

# ======================================================================================
# 5. TD-space scatter (pointwise) next to raw
# ======================================================================================
fig, ax = plt.subplots(1, 2, figsize=(12, 5.6))
for a_, (P, T, ttl, lim) in zip(ax, [
        (sens_p, sens_t, "Raw sensitivity (dB)", (0, 38)),
        (td_p, td_t, "Total deviation (dB)", (-38, 12))]):
    m = ~np.isnan(T)
    a_.hexbin(T[m], P[m], gridsize=45, cmap="viridis", mincnt=1, bins="log")
    lo, hi = lim
    a_.plot([lo, hi], [lo, hi], "w--", lw=1, alpha=.7)
    st = pooled(P, T)
    b, i0 = np.polyfit(T[m], P[m], 1)
    xs = np.array([lo, hi]); a_.plot(xs, b * xs + i0, "r-", lw=1.6,
                                     label=f"fit y={b:.2f}x+{i0:.1f}")
    a_.set_xlim(lim); a_.set_ylim(lim); a_.set_aspect("equal")
    a_.set_xlabel(f"True {ttl}"); a_.set_ylabel(f"Pred {ttl}")
    a_.set_title(f"{ttl.split(' (')[0]}\nMAE {st['mae']:.2f}  r {st['r']:.3f}  slope {st['slope']:.2f}")
    a_.legend(loc="upper left", fontsize=9)
fig.suptitle(f"{TAG}: pooled OOF (631 eyes) — raw sensitivity vs total deviation", fontsize=12)
fig.tight_layout()
out = os.path.join(AUTO, f"{TAG}_td_scatter.png")
fig.savefig(out, dpi=130); print(f"\n[5] scatter -> {out}")

# machine-readable dump
json.dump({"tag": TAG, "n_eyes": len(md_t), "pointwise": {"raw_sensitivity": S, "total_deviation": D},
           "per_eye": {"mean_sensitivity": MDs, "mean_deviation": MDd}},
          open(os.path.join(AUTO, f"{TAG}_td_report.json"), "w"), indent=2)
print(f"    json -> {os.path.join(AUTO, TAG + '_td_report.json')}")
